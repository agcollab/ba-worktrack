"""Turn a table of rows into a downloadable CSV, Excel or PDF file.

Every export (tasks, timesheets, requirements, reports) builds an ExportTable
and hands it to one of the writers below, so all formats stay consistent.
"""
import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.config import settings

FONT_DIR = Path(__file__).parent / "fonts"
BRAND = (79, 70, 229)  # indigo, matches the UI
MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
# Spreadsheet apps run cells that start with these as formulas (CSV injection)
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


@dataclass
class ExportTable:
    title: str
    columns: list[str]
    rows: list[list[Any]]
    subtitle: str | None = None
    # Relative column widths for the PDF (defaults to equal); bigger = wider
    pdf_widths: list[int] | None = None
    sheet_name: str = "Data"
    totals: list[Any] | None = None  # optional final row


@dataclass
class ReportSection:
    """A titled block in a multi-section PDF / workbook (used by the weekly report)."""
    heading: str
    table: ExportTable | None = None
    text: str | None = None
    kpis: list[tuple[str, str]] = field(default_factory=list)


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return round(value, 2)
    return value if isinstance(value, int | float) else str(value)


def _safe_text(value: Any) -> Any:
    value = _cell(value)
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


# ---- CSV --------------------------------------------------------------------

def to_csv(table: ExportTable) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(table.columns)
    for row in table.rows:
        writer.writerow([_safe_text(v) for v in row])
    if table.totals:
        writer.writerow([_safe_text(v) for v in table.totals])
    # BOM so Excel opens UTF-8 (names with accents) correctly
    return ("﻿" + buffer.getvalue()).encode("utf-8")


# ---- Excel ------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", start_color="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_sheet(ws: Worksheet, table: ExportTable, start_row: int = 1) -> int:
    """Write one table at start_row; returns the next free row."""
    row_no = start_row
    if table.title:
        ws.cell(row=row_no, column=1, value=table.title).font = Font(bold=True, size=14)
        row_no += 1
    if table.subtitle:
        ws.cell(row=row_no, column=1, value=table.subtitle).font = Font(italic=True, color="666666")
        row_no += 1
    if table.title or table.subtitle:
        row_no += 1
    header_row = row_no
    for col, name in enumerate(table.columns, start=1):
        cell = ws.cell(row=row_no, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for values in table.rows:
        row_no += 1
        for col, value in enumerate(values, start=1):
            v = value if isinstance(value, date | datetime) else _safe_text(value)
            cell = ws.cell(row=row_no, column=col, value=v)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(v, str) and "\n" in v:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    if table.totals:
        row_no += 1
        for col, value in enumerate(table.totals, start=1):
            ws.cell(row=row_no, column=col, value=_safe_text(value)).font = Font(bold=True)
    # Column widths from the longest value (capped so long descriptions wrap)
    for col, name in enumerate(table.columns, start=1):
        longest = max([len(str(name))] + [len(str(_cell(r[col - 1]))) for r in table.rows[:500] if col - 1 < len(r)])
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, min(60, longest + 2))
    if start_row == 1 and table.rows:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(table.columns))}{header_row + len(table.rows)}"
    return row_no + 2


def to_xlsx(table: ExportTable) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = table.sheet_name[:31]
    _write_sheet(ws, table)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sections_to_xlsx(title: str, sections: list[ReportSection]) -> bytes:
    """Summary sheet with KPIs and text, then one sheet per table."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    row = 3
    for section in sections:
        if section.kpis or section.text:
            summary.cell(row=row, column=1, value=section.heading).font = Font(bold=True, size=12)
            row += 1
            for label, value in section.kpis:
                summary.cell(row=row, column=1, value=label)
                summary.cell(row=row, column=2, value=value).font = Font(bold=True)
                row += 1
            if section.text:
                cell = summary.cell(row=row, column=1, value=section.text)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
                summary.row_dimensions[row].height = max(30, 15 * (len(section.text) // 90 + 1))
                row += 1
            row += 1
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 18
    used = {"Summary"}
    for section in sections:
        if section.table is None:
            continue
        name = section.table.sheet_name[:31]
        while name in used:
            name = name[:28] + str(len(used))
        used.add(name)
        _write_sheet(wb.create_sheet(name), section.table)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---- PDF --------------------------------------------------------------------

class BrandedPDF(FPDF):
    def __init__(self, title: str, subtitle: str | None = None, orientation: str = "L"):
        super().__init__(orientation=orientation, unit="mm", format="A4")
        self.doc_title = title
        self.doc_subtitle = subtitle
        self.add_font("DejaVu", "", str(FONT_DIR / "DejaVuSansCondensed.ttf"))
        self.add_font("DejaVu", "B", str(FONT_DIR / "DejaVuSansCondensed-Bold.ttf"))
        self.set_auto_page_break(auto=True, margin=15)
        self.set_margins(12, 12, 12)
        self.set_title(title)
        self.set_creator(settings.APP_NAME)
        self.alias_nb_pages()

    def header(self) -> None:
        self.set_fill_color(*BRAND)
        self.rect(0, 0, self.w, 3, "F")
        self.set_y(8)
        self.set_font("DejaVu", "B", 15)
        self.set_text_color(17, 24, 39)
        self.cell(0, 8, self.doc_title, new_x="LMARGIN", new_y="NEXT")
        if self.doc_subtitle:
            self.set_font("DejaVu", "", 9)
            self.set_text_color(100, 100, 110)
            self.cell(0, 5, self.doc_subtitle, new_x="LMARGIN", new_y="NEXT")
        self.ln(3)
        self.set_text_color(17, 24, 39)

    def footer(self) -> None:
        self.set_y(-10)
        self.set_font("DejaVu", "", 8)
        self.set_text_color(130, 130, 140)
        self.cell(0, 5, f"{settings.APP_NAME} · generated {datetime.now():%d %b %Y %H:%M}", align="L")
        self.cell(0, 5, f"Page {self.page_no()} of {{nb}}", align="R")

    def heading(self, text: str) -> None:
        self.ln(2)
        self.set_font("DejaVu", "B", 12)
        self.set_text_color(*BRAND)
        self.cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(17, 24, 39)

    def kpi_row(self, kpis: list[tuple[str, str]]) -> None:
        if not kpis:
            return
        per_row = 4
        width = (self.w - self.l_margin - self.r_margin - (per_row - 1) * 3) / per_row
        for start in range(0, len(kpis), per_row):
            y = self.get_y()
            for i, (label, value) in enumerate(kpis[start:start + per_row]):
                x = self.l_margin + i * (width + 3)
                self.set_fill_color(243, 244, 250)
                self.rect(x, y, width, 16, "F")
                self.set_xy(x + 3, y + 2)
                self.set_font("DejaVu", "", 8)
                self.set_text_color(100, 100, 110)
                self.cell(width - 6, 4, label)
                self.set_xy(x + 3, y + 7)
                self.set_font("DejaVu", "B", 13)
                self.set_text_color(17, 24, 39)
                self.cell(width - 6, 7, value)
            self.set_y(y + 19)

    def paragraph(self, text: str) -> None:
        self.set_font("DejaVu", "", 9.5)
        self.multi_cell(0, 5, text, new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def data_table(self, table: ExportTable) -> None:
        if not table.rows:
            self.set_font("DejaVu", "", 9)
            self.set_text_color(120, 120, 130)
            self.cell(0, 6, "No data for this period.", new_x="LMARGIN", new_y="NEXT")
            self.set_text_color(17, 24, 39)
            return
        self.set_font("DejaVu", "", 8)
        widths = table.pdf_widths or [1] * len(table.columns)
        heading_style = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=BRAND)
        with self.table(
            col_widths=widths,
            headings_style=heading_style,
            line_height=4.6,
            cell_fill_color=(247, 247, 252),
            cell_fill_mode="ROWS",
            borders_layout="HORIZONTAL_LINES",
            padding=1.2,
        ) as pdf_table:
            pdf_table.row([str(c) for c in table.columns])
            for values in table.rows:
                pdf_table.row([_pdf_text(v) for v in values])
            if table.totals:
                pdf_table.row([_pdf_text(v) for v in table.totals], style=FontFace(emphasis="BOLD"))
        self.ln(2)


def _pdf_text(value: Any, limit: int = 600) -> str:
    text = str(_cell(value))
    return text if len(text) <= limit else text[: limit - 1] + "…"


def to_pdf(table: ExportTable) -> bytes:
    pdf = BrandedPDF(table.title, table.subtitle)
    pdf.add_page()
    pdf.data_table(table)
    return bytes(pdf.output())


def sections_to_pdf(title: str, subtitle: str | None, sections: list[ReportSection]) -> bytes:
    pdf = BrandedPDF(title, subtitle)
    pdf.add_page()
    for section in sections:
        pdf.heading(section.heading)
        pdf.kpi_row(section.kpis)
        if section.text:
            pdf.paragraph(section.text)
        if section.table is not None:
            pdf.data_table(section.table)
    return bytes(pdf.output())


def render(table: ExportTable, fmt: str) -> bytes:
    if fmt == "csv":
        return to_csv(table)
    if fmt == "xlsx":
        return to_xlsx(table)
    if fmt == "pdf":
        return to_pdf(table)
    raise ValueError(f"Unknown export format: {fmt}")


def filename(base: str, fmt: str) -> str:
    return f"{base}-{date.today():%Y%m%d}.{fmt}"
