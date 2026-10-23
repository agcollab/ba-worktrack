"""CSV / Excel / PDF exports of tasks, time, requirements and reports."""
import csv
import io
from datetime import date

from openpyxl import load_workbook

from app.reports.exporters import ExportTable, to_csv
from app.services import requirements as requirement_service
from app.services import tasks as task_service
from app.services import time_tracking
from tests.conftest import login


def seed_some(db, people):
    t1 = task_service.create_task(db, people["pm"], people["pay"].id, "Refund UAC", assignee_id=people["asha"].id,
                                  estimate_hours=4)
    task_service.create_task(db, people["pm"], people["crm"].id, "Portal login stories")
    time_tracking.log_time(db, people["asha"], t1, date.today(), 2)
    req = requirement_service.save_requirement(db, people["pm"], people["pay"].id, "Refunds over 500 need approval",
                                               None, task_id=t1.id)
    requirement_service.add_criterion(db, req, "a refund of 600", "the agent submits it", "it waits for approval")
    requirement_service.add_test_case(db, req, "Approval needed", "Create refund of 600\nSubmit", "Status is Pending")
    db.commit()
    return t1


def test_tasks_csv_respects_filters_and_visibility(client, db, people):
    seed_some(db, people)
    login(client, "asha@test.com")
    response = client.get("/exports/tasks.csv")
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0][:2] == ["Key", "Title"]
    titles = [r[1] for r in rows[1:-1]]  # last row is the totals
    assert titles == ["Refund UAC"]  # Asha can't see the CRM task


def test_tasks_xlsx_and_pdf(client, db, people):
    seed_some(db, people)
    login(client, "pm@test.com")
    xlsx = client.get(f"/exports/tasks.xlsx?project_id={people['pay'].id}")
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    values = [c.value for c in sheet[4]]  # title, subtitle, blank, header
    assert values[:2] == ["Key", "Title"]
    assert sheet.cell(row=5, column=2).value == "Refund UAC"
    pdf = client.get("/exports/tasks.pdf")
    assert pdf.content.startswith(b"%PDF")


def test_time_requirements_and_report_exports(client, db, people):
    seed_some(db, people)
    login(client, "pm@test.com")
    assert b"Refund UAC" in client.get("/exports/time.csv").content
    req = load_workbook(io.BytesIO(client.get("/exports/requirements.xlsx").content)).active
    assert any("Given a refund of 600" in str(c.value) for row in req.iter_rows() for c in row)
    assert b"Approval needed" in client.get("/exports/test-cases.csv").content
    assert client.get("/exports/report.pdf").content.startswith(b"%PDF")
    report = load_workbook(io.BytesIO(client.get("/exports/report.xlsx").content))
    assert report.sheetnames[0] == "Summary" and "Hours" in report.sheetnames


def test_unknown_format_is_404(client, people):
    login(client, "pm@test.com")
    assert client.get("/exports/tasks.docx").status_code == 404


def test_csv_formula_injection_is_neutralised():
    data = to_csv(ExportTable(title="t", columns=["a"], rows=[["=HYPERLINK(\"http://x\")"], ["+1"], ["ok"]]))
    text = data.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'+1" in text and "\nok" in text.replace("\r", "")


def test_exports_need_login(client):
    assert client.get("/exports/tasks.csv", follow_redirects=False).status_code == 303
